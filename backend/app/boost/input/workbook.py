"""Spec-independent workbook reader: cell text, merged ranges, column widths and images.

Images are read from the drawing parts directly so that both the anchor cell and the
displayed size (as the user sized the picture in Excel) are available.
"""
from __future__ import annotations

import posixpath
import re
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime, time
from pathlib import Path

from lxml import etree
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter

from .model import ERROR, WARNING, Issue

MAX_UNCOMPRESSED_BYTES = 200 * 1024 * 1024
EMU_PER_PX = 9525
OLE_SIGNATURE = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"

XNS = {
    "main": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "pr": "http://schemas.openxmlformats.org/package/2006/relationships",
    "xdr": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "mc": "http://schemas.openxmlformats.org/markup-compatibility/2006",
}
CONTROL_CHARACTERS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
FIXED_DECIMALS = re.compile(r"^0(?:\.(0+))?(?:_.?)?\s*$")
PERCENT = re.compile(r"^0(?:\.(0+))?%$")


class WorkbookError(ValueError):
    def __init__(self, code: str, message: str):
        super().__init__(message)
        self.code = code


@dataclass(frozen=True, slots=True)
class SheetImage:
    row: int
    col: int
    data: bytes
    name: str
    display_w: int
    display_h: int

    @property
    def coord(self) -> str:
        return f"{get_column_letter(self.col)}{self.row}"


@dataclass(slots=True)
class SheetGrid:
    name: str
    hidden: bool = False
    cells: dict[tuple[int, int], str] = field(default_factory=dict)
    merges: list[tuple[int, int, int, int]] = field(default_factory=list)  # min_row, min_col, max_row, max_col
    images: list[SheetImage] = field(default_factory=list)
    col_widths: dict[int, float] = field(default_factory=dict)
    max_row: int = 0
    _covering: dict[tuple[int, int], tuple[int, int, int, int]] = field(default_factory=dict, init=False, repr=False)

    def __post_init__(self) -> None:
        for merge in self.merges:
            min_row, min_col, max_row, max_col = merge
            for row in range(min_row, max_row + 1):
                for col in range(min_col, max_col + 1):
                    self._covering[(row, col)] = merge

    def merge_at(self, row: int, col: int) -> tuple[int, int, int, int] | None:
        return self._covering.get((row, col))

    def is_covered(self, row: int, col: int) -> bool:
        merge = self.merge_at(row, col)
        return merge is not None and (merge[0], merge[1]) != (row, col)

    def text(self, row: int, col: int) -> str:
        return self.cells.get((row, col), "")

    def anchor_text(self, row: int, col: int) -> str:
        merge = self.merge_at(row, col)
        return self.text(merge[0], merge[1]) if merge else self.text(row, col)

    def width(self, col: int) -> float:
        return self.col_widths.get(col, 8.43)


def coord(row: int, col: int) -> str:
    return f"{get_column_letter(col)}{row}"


def _clean(value: str) -> str:
    value = value.replace("_x000D_", "").replace("\r\n", "\n").replace("\r", "\n")
    return CONTROL_CHARACTERS.sub("", value).strip()


def _general(value: float) -> str:
    if float(value).is_integer():
        return str(int(value))
    return format(value, ".15g")


def format_value(value, number_format: str) -> tuple[str, bool]:
    """Return (display text, supported) for a cell value."""
    if value is None:
        return "", True
    if isinstance(value, str):
        return _clean(value), True
    if isinstance(value, bool):
        return ("TRUE" if value else "FALSE"), True
    if isinstance(value, datetime):
        return (value.date().isoformat() if value.time() == time(0) else value.isoformat(sep=" ", timespec="minutes")), True
    if isinstance(value, date):
        return value.isoformat(), True
    if isinstance(value, int | float):
        fmt = (number_format or "General").strip()
        if fmt in {"General", "@"}:
            return _general(value), True
        percent = PERCENT.match(fmt)
        if percent:
            decimals = len(percent.group(1) or "")
            return f"{value * 100:.{decimals}f}%", True
        fixed = FIXED_DECIMALS.match(fmt)
        if fixed:
            decimals = len(fixed.group(1) or "")
            return f"{value:.{decimals}f}", True
        return _general(value), False
    return _clean(str(value)), True


def _resolve(base: str, target: str) -> str:
    if target.startswith("/"):
        return target.lstrip("/")
    return posixpath.normpath(posixpath.join(posixpath.dirname(base), target))


def _rels(archive: zipfile.ZipFile, part: str) -> dict[str, tuple[str, str]]:
    rels_name = posixpath.join(posixpath.dirname(part), "_rels", posixpath.basename(part) + ".rels")
    if rels_name not in archive.namelist():
        return {}
    root = etree.fromstring(archive.read(rels_name))
    return {
        rel.get("Id"): (rel.get("Type", ""), _resolve(part, rel.get("Target", "")))
        for rel in root.findall("pr:Relationship", XNS)
        if rel.get("TargetMode") != "External"
    }


def _sheet_parts(archive: zipfile.ZipFile) -> dict[str, str]:
    workbook = etree.fromstring(archive.read("xl/workbook.xml"))
    rels = _rels(archive, "xl/workbook.xml")
    parts = {}
    for sheet in workbook.findall("main:sheets/main:sheet", XNS):
        rel = rels.get(sheet.get(f"{{{XNS['r']}}}id"))
        if rel:
            parts[sheet.get("name")] = rel[1]
    return parts


def _drawing_images(archive: zipfile.ZipFile, sheet_part: str, issues: list[Issue], sheet_name: str) -> list[SheetImage]:
    images: list[SheetImage] = []
    for rel_type, drawing_part in _rels(archive, sheet_part).values():
        if not rel_type.endswith("/drawing") or drawing_part not in archive.namelist():
            continue
        drawing = etree.fromstring(archive.read(drawing_part))
        media = _rels(archive, drawing_part)
        for picture in drawing.xpath(".//xdr:pic[not(ancestor::mc:Fallback)]", namespaces=XNS):
            anchor = next(
                (node for node in picture.iterancestors() if etree.QName(node).localname in {"twoCellAnchor", "oneCellAnchor", "absoluteAnchor"}),
                None,
            )
            start = anchor.find("xdr:from", XNS) if anchor is not None else None
            if start is None:
                issues.append(Issue(WARNING, "IMAGE_NOT_ANCHORED", "图片没有锚定到单元格，已放在工作表开头", sheet_name, "A1"))
                row, col = 1, 1
            else:
                row = int(start.findtext("xdr:row", "0", XNS)) + 1
                col = int(start.findtext("xdr:col", "0", XNS)) + 1
            embed = picture.xpath("string(xdr:blipFill/a:blip/@r:embed)", namespaces=XNS)
            target = media.get(embed)
            if target is None or target[1] not in archive.namelist():
                issues.append(Issue(ERROR, "IMAGE_UNREADABLE", "图片数据缺失，无法读取", sheet_name, coord(row, col)))
                continue
            extent = picture.find("xdr:spPr/a:xfrm/a:ext", XNS)
            if extent is None and anchor is not None:
                extent = anchor.find("xdr:ext", XNS)
            width = round(int(extent.get("cx", "0")) / EMU_PER_PX) if extent is not None else 0
            height = round(int(extent.get("cy", "0")) / EMU_PER_PX) if extent is not None else 0
            images.append(SheetImage(row, col, archive.read(target[1]), posixpath.basename(target[1]), width, height))
    images.sort(key=lambda image: (image.row, image.col))
    return images


def read_workbook(path: Path) -> tuple[list[SheetGrid], list[Issue]]:
    head = path.read_bytes()[:8]
    if head == OLE_SIGNATURE:
        raise WorkbookError("XLS_NOT_SUPPORTED", "不支持旧版 .xls 文件，请在 Excel 中另存为 .xlsx 后上传")
    if not zipfile.is_zipfile(path):
        raise WorkbookError("WORKBOOK_UNREADABLE", "文件不是有效的 .xlsx 工作簿")
    issues: list[Issue] = []
    with zipfile.ZipFile(path) as archive:
        if sum(info.file_size for info in archive.infolist()) > MAX_UNCOMPRESSED_BYTES:
            raise WorkbookError("WORKBOOK_TOO_LARGE", "工作簿解压后体积过大")
        if "xl/workbook.xml" not in archive.namelist():
            raise WorkbookError("WORKBOOK_UNREADABLE", "文件不是有效的 .xlsx 工作簿")
        sheet_parts = _sheet_parts(archive)
        try:
            workbook = load_workbook(path, data_only=True)
        except Exception as exc:
            raise WorkbookError("WORKBOOK_UNREADABLE", f"无法读取工作簿：{exc}") from exc
        grids: list[SheetGrid] = []
        for worksheet in workbook.worksheets:
            cells: dict[tuple[int, int], str] = {}
            for row in worksheet.iter_rows():
                for cell in row:
                    if cell.value is None:
                        continue
                    text, supported = format_value(cell.value, cell.number_format)
                    if not supported:
                        issues.append(Issue(WARNING, "UNSUPPORTED_NUMBER_FORMAT",
                                            f"数字格式「{cell.number_format}」按常规格式输出：{text}", worksheet.title, cell.coordinate))
                    if text:
                        cells[(cell.row, cell.column)] = text
            merges = [(r.min_row, r.min_col, r.max_row, r.max_col) for r in worksheet.merged_cells.ranges]
            widths: dict[int, float] = {}
            for letter, dimension in worksheet.column_dimensions.items():
                if not dimension.customWidth or not dimension.width:
                    continue
                first = dimension.min or column_index_from_string(letter)
                for index in range(first, (dimension.max or first) + 1):
                    widths[index] = float(dimension.width)
            part = sheet_parts.get(worksheet.title)
            images = _drawing_images(archive, part, issues, worksheet.title) if part else []
            max_row = max([row for row, _ in cells] + [merge[2] for merge in merges] + [image.row for image in images] + [0])
            grids.append(SheetGrid(worksheet.title, worksheet.sheet_state != "visible", cells, merges, images, widths, max_row))
        workbook.close()
    return grids, issues
