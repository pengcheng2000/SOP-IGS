from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from PIL import Image


def build_sample_workbook(path: Path) -> Path:
    """A small workbook following the boost input spec (draft v1, format column J)."""
    images = path.parent / "fixture-images"
    images.mkdir(parents=True, exist_ok=True)
    icon = images / "icon.png"
    figure = images / "figure.png"
    Image.new("RGB", (40, 20), (200, 30, 30)).save(icon)
    Image.new("RGB", (600, 300), (30, 120, 60)).save(figure)

    workbook = Workbook()
    appearance = workbook.active
    appearance.title = "外观检查"
    appearance["A1"] = "集装箱尺寸检查"
    appearance.merge_cells("A1:A2")
    appearance["B1"] = "尺寸符合要求"
    appearance["B2"] = "铭牌正确"
    appearance["J2"] = "注意"
    appearance["A3"] = "箱门检查"
    appearance["B3"] = "开闭灵活"
    appearance["J3"] = "列表"
    appearance["B4"] = "密封良好"
    appearance["J4"] = "列表"
    appearance["B5"] = "子项"
    appearance["J5"] = "子列表"
    appearance["B6"] = "GB项目"
    appearance["C6"] = "L1（黄）"

    insulation = workbook.create_sheet("绝缘电阻与耐压测试")
    insulation["A1"] = "测控电源柜"
    insulation["B1"] = "测试条件"
    insulation["B2"] = "测试部位"
    insulation.merge_cells("B2:B3")
    insulation["C2"] = "要求"
    insulation.merge_cells("C2:D2")
    insulation["J2"] = "表头"
    insulation["C3"] = "绝缘电阻"
    insulation["D3"] = "耐压测试"
    insulation["J3"] = "表头"
    insulation["B4"] = "短接对地"
    insulation["C4"] = "≥1MΩ"
    insulation["D4"] = 1.01
    insulation["D4"].number_format = "0%"
    insulation["J4"] = "表格"
    insulation["B5"] = "急停信号跳闸测试"
    insulation["J5"] = "四级标题"
    insulation["B6"] = "按下急停按钮"
    insulation["J6"] = "列表"
    insulation["C7"] = "0002"
    insulation["B7"] = "综保密码"
    insulation.column_dimensions["B"].width = 30
    insulation.add_image(XLImage(str(icon)), "C4")
    insulation.add_image(XLImage(str(figure)), "B8")

    revisions = workbook.create_sheet("版本修改记录")
    revisions["A1"] = "2026年10月14日，A0"
    instructions = workbook.create_sheet("填写说明")
    instructions["A1"] = "此页不会生成"
    workbook.save(path)
    return path


@pytest.fixture
def sample_workbook(tmp_path: Path) -> Path:
    return build_sample_workbook(tmp_path / "boost-input.xlsx")
