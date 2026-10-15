from pathlib import Path

import pytest
from app.boost.input import parse_workbook
from app.boost.input.model import Heading, ListBlock, Paragraph, Table
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
TEMPLATE = ROOT / "templates" / "boost-integrated-unit" / "v1" / "input-template.xlsx"


def parse(path: Path, tmp_path: Path):
    return parse_workbook(path, tmp_path / "images")


def test_standard_template_parses_without_errors(tmp_path: Path) -> None:
    result = parse(TEMPLATE, tmp_path)
    assert result.valid, [issue for issue in result.issues if issue.severity == "error"]
    assert result.summary()["chapter_count"] == 11
    assert result.summary()["table_count"] >= 15 and result.summary()["image_count"] == 16
    titles = [chapter.title for chapter in result.document.chapters]
    assert titles[0] == "外观检查" and "系统供电测试" in titles and "填写说明" not in titles
    assert result.document.revision_history  # 版本修改记录 sheet


def test_template_reproduces_the_case_structure(tmp_path: Path) -> None:
    document = parse(TEMPLATE, tmp_path).document
    chapters = {chapter.title: chapter.items for chapter in document.chapters}

    # no 三级标题 for chapters whose text sits directly under the chapter
    assert not any(isinstance(item, Heading) for item in chapters["防水测试"] + chapters["接地连续性"])
    assert [type(item).__name__ for item in chapters["防水测试"]] == ["Paragraph", "ListBlock"]

    # 5.3.3: numbered list whose last item has lettered sub-items for the colour standards
    wiring = [item for item in chapters["接线正确性"] if isinstance(item, ListBlock)]
    assert [block.style for block in wiring] == ["decimal", "letter"]

    # 800V/380V/220V/24V loops: one numbered list kept going across indented paragraphs
    insulation = chapters["绝缘电阻与耐压测试"]
    first_list = next(item for item in insulation if isinstance(item, ListBlock))
    assert [entry.text for entry in first_list.items][:1] == ["800V回路"]
    assert sum(isinstance(item, Table) for item in insulation) == 5
    assert any(isinstance(item, Paragraph) and item.kind == "indent" for item in insulation)

    power = chapters["系统供电测试"]
    level4 = [item.text for item in power if isinstance(item, Heading) and item.level == 4]
    assert {"内部通讯测试", "外部通讯线路测试", "急停信号跳闸测试", "变压器压力异常跳闸测试"} <= set(level4)
    assert any(isinstance(item, Heading) and item.level == 3 and item.text == "PCS输入干接点测试" for item in power)


def test_template_has_instructions_and_validation(tmp_path: Path) -> None:
    workbook = load_workbook(TEMPLATE)
    assert workbook.sheetnames[0] == "填写说明"
    text = "\n".join(str(cell.value) for row in workbook["填写说明"].iter_rows() for cell in row if cell.value)
    for name in ("缩进段落", "字母列表", "四级标题", "表头", "表格", "A 列", "J 列"):
        assert name in text
    sheet = workbook["外观检查"]
    assert sheet.data_validations.dataValidation and sheet["J1"].comment is not None


def test_workbook_without_any_format_marker_is_rejected(tmp_path: Path) -> None:
    workbook = load_workbook(TEMPLATE)
    for sheet in workbook.worksheets:
        for row in range(1, sheet.max_row + 1):
            sheet.cell(row, 10).value = None
    legacy = tmp_path / "legacy.xlsx"
    workbook.save(legacy)
    result = parse(legacy, tmp_path)
    assert not result.valid
    assert [issue.code for issue in result.issues if issue.severity == "error"] == ["NO_FORMAT_COLUMN"]
    assert "标准模板" in next(i.message for i in result.issues if i.code == "NO_FORMAT_COLUMN")


@pytest.mark.skipif(
    not (ROOT / "SOP_case" / "升压一体机案例" / "升压一体机测试项分类表.xlsx").is_file(),
    reason="case workbook is project data and not in git",
)
def test_original_case_workbook_is_rejected_with_guidance(tmp_path: Path) -> None:
    result = parse(ROOT / "SOP_case" / "升压一体机案例" / "升压一体机测试项分类表.xlsx", tmp_path)
    assert not result.valid and "NO_FORMAT_COLUMN" in [issue.code for issue in result.issues]
