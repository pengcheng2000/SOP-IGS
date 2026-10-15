"""Build the standard 升压一体机 input template (input-template.xlsx) from the case documents.

    python tools/build_boost_input_template.py

The case workbook (values and wording) is annotated with a 格式 in column J for every content row; the
formats are derived from how the case Word document lays the same text out (see boost_template_derive.py),
plus the explicit decisions below for parts that exist only in the Excel. A 填写说明 sheet is added.
Needs the case files in SOP_case/ (not in git); the produced template is committed.
"""
from __future__ import annotations

import argparse
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "tools"))
sys.path.insert(0, str(ROOT / "backend"))

import boost_template_derive as derive
from app.boost.input.workbook import read_workbook

FORMAT_COLUMN = 10
FORMAT_FILL = PatternFill("solid", fgColor="FFF2CC")
FORMATS = [
    ("段落", "普通正文，首行缩进两个字符。", "升压一体机符合尺寸：长*宽*高…"),
    ("缩进段落", "挂在上一个列表项下面的说明文字（缩进更多），不会打断列表编号。", "列表项「800V回路：」下面的各条测试电压"),
    ("列表", "编号列表 1) 2) 3)，连续的「列表」行是同一个列表；图片或缩进段落不会打断编号。", "按照“布线表”核对各柜的线缆敷设正确"),
    ("字母列表", "字母列表 a. b. c.", "GB项目：L1（黄）…"),
    ("符号列表", "圆点列表 ●", "上电前确认断路器…"),
    ("子列表", "挂在上一个列表项下的二级列表 a) b)。", ""),
    ("四级标题", "5.x.y.z 小标题（需要先有 A 列三级标题）。", "内部通讯测试、急停信号跳闸测试"),
    ("注意", "以 ★ 开头的注意事项（★ 自动添加）。", "变压器高压侧不做耐压测试。"),
    ("表头", "表格的表头行，可以有多行；合并单元格会按原样合并。", "测试部位 / 要求 / 绝缘电阻 / 耐压测试"),
    ("表格", "表格的数据行。连续的「表头」「表格」行组成一张表。", ""),
]

# Rows that exist only in the Excel (the case Word lacks them): sheet -> {row: 格式}.
EXCEL_ONLY: dict[str, dict[int, str]] = {
    "系统供电测试": {
        31: "表头", **{r: "表格" for r in range(32, 38)},
        47: "表头", **{r: "表格" for r in range(48, 57)},
        57: "列表", 58: "列表",
        62: "段落", 63: "表头", **{r: "表格" for r in range(64, 68)},
        89: "段落",
        238: "四级标题+段落+段落",
        242: "表头", **{r: "表格" for r in range(243, 259)},
        260: "列表+缩进段落", 261: "列表+缩进段落+缩进段落", 262: "缩进段落",
        **{r: "列表+缩进段落" for r in (263, 264, 265, 266, 269)},
        267: "缩进段落", 268: "缩进段落", **{r: "缩进段落" for r in (270, 271, 272, 273)},
        274: "四级标题+段落", 275: "表头", **{r: "表格" for r in range(276, 303)},
        303: "段落", 304: "段落",
    },
}


def fix_cells(workbook) -> None:
    """Repairs that make the case workbook a valid standard (each follows the case Word document)."""
    insulation = workbook["绝缘电阻与耐压测试"]
    if not insulation["C24"].value:  # the AC-side table lacks its 要求 header
        insulation["C24"] = "要求"
        insulation.merge_cells("C24:D24")

    versions = workbook["软件版本"]  # the standard keeps column A for headings, so the table moves one column right
    for merged in list(versions.merged_cells.ranges):
        versions.unmerge_cells(str(merged))
    versions.move_range("A1:B4", cols=1)
    versions.merge_cells("B2:B3")

    polarity = workbook["互感器极性测试"]  # the two checks are list items with lettered sub-items, not headings
    cells = {row: polarity.cell(row, 2).value for row in (12, 13, 14, 15)}
    labels = {12: polarity["A12"].value, 14: polarity["A14"].value}
    for merged in list(polarity.merged_cells.ranges):
        if str(merged) in {"A12:A13", "A14:A15"}:
            polarity.unmerge_cells(str(merged))
    for row in range(12, 18):
        polarity.cell(row, 1).value = polarity.cell(row, 2).value = None
    layout = {12: labels[12] + "：", 13: cells[12], 14: cells[13], 15: labels[14] + "：", 16: cells[14], 17: cells[15]}
    for row, text in layout.items():
        polarity.cell(row, 2, text)
    for row in layout:  # the rows were laid out for the case's merged cells: give them one plain style
        polarity.row_dimensions[row].height = None
        for column in range(1, 3):
            cell = polarity.cell(row, column)
            cell.border = Border()
            cell.font = Font(name="宋体", size=11)
            cell.alignment = Alignment(wrap_text=True, vertical="center")

    resistance = workbook["变压器直流连接电阻测试"]
    resistance["A3"] = resistance["A3"].value.replace("试过程中", "测试过程中", 1) if str(resistance["A3"].value).startswith("试") else resistance["A3"].value

    power = workbook["系统供电测试"]
    for row in range(70, 86):  # signal addresses are shown with one decimal (0.0, 0.1 … 1.7)
        power.cell(row, 3).number_format = "0.0"
    power["B89"] = "软件配置（艾莫迅）"  # no manual line break inside a paragraph
    power["A239"] = "充气柜设备测试（施耐德）"
    power["A260"] = "充气柜设备测试（正泰）"
    # 「PCS输入干接点测试」 is its own section in the case Word; in the workbook it sat inside 通讯测试 as a label.
    for merged in list(power.merged_cells.ranges):
        if str(merged) in {"A172:A228", "D211:D228"}:
            power.unmerge_cells(str(merged))
    power.merge_cells("A172:A210")
    power["A211"] = "PCS输入干接点测试"
    power.merge_cells("A211:A228")
    power["D211"] = None


def style_and_validate(sheet) -> None:
    last = max(sheet.max_row, 1) + 40
    validation = DataValidation(
        type="list",
        formula1='"' + ",".join(name for name, _, _ in FORMATS) + '"',
        allow_blank=True,
        showErrorMessage=False,  # combinations such as 四级标题+列表 are typed by hand
        promptTitle="格式",
        prompt="选择这一行的格式；一行里有多个单元格要分别成条目时，用 + 连接，例如：四级标题+列表+缩进段落",
        showInputMessage=True,
    )
    sheet.add_data_validation(validation)
    validation.add(f"J1:J{last}")
    sheet.column_dimensions["J"].width = 22
    for row in range(1, sheet.max_row + 1):
        cell = sheet.cell(row, FORMAT_COLUMN)
        cell.fill = FORMAT_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    sheet["J1"].comment = Comment(
        "J 列「格式」：每个有内容的行都要填写。\n点单元格从下拉里选；一行有多个条目时用 + 连接。\n详见「填写说明」工作表。", "SOP"
    )


def add_instructions(workbook) -> None:
    sheet = workbook.create_sheet("填写说明", 0)
    sheet.sheet_properties.tabColor = "0F766E"
    sheet.column_dimensions["A"].width = 14
    sheet.column_dimensions["B"].width = 58
    sheet.column_dimensions["C"].width = 46
    bold = Font(bold=True)
    title = Font(bold=True, size=14)
    wrap = Alignment(wrap_text=True, vertical="top")
    rows: list[tuple[str, str, str]] = [
        ("升压一体机测试项分类表 · 填写说明", "", ""),
        ("", "", ""),
        ("整体规则", "", ""),
        ("工作表", "一张工作表 = 报告第 5 章下的一个章节：工作表名就是章节标题，工作表的先后顺序就是章节顺序（5.2、5.3…）。「版本修改记录」工作表每行写一条版本记录。「填写说明」不会生成到文档里。", ""),
        ("A 列", "三级标题（5.x.y）。同一个标题下有多行内容时，可以把 A 列纵向合并覆盖这些行。A 列只写简短的标题；不需要三级标题的章节，A 列留空。", ""),
        ("B–I 列", "内容。正文写在 B 列，需要分栏的内容（表格、「标签 + 说明」）按列写在 B–I。", ""),
        ("J 列 格式", "每个有内容的行都要在 J 列选一个格式（下拉选择）。J 列留空 = 普通段落。没有任何格式标记的表格会被拒绝上传。", ""),
        ("", "", ""),
        ("格式一览", "含义", "例子"),
        *[(name, meaning, example) for name, meaning, example in FORMATS],
        ("", "", ""),
        ("进阶写法", "", ""),
        ("一行多个条目", "一行里的内容单元格要分别成为不同条目时，格式用 + 连接，依次对应这一行从左到右的内容单元格。例如 B 列「内部通讯测试」、D 列「模块数据与 EMS 一致」，格式写「四级标题+列表」，就生成一个四级标题和一个列表项。格式个数少于单元格个数时，最后一个格式吃掉剩余的全部单元格（用「；」连接）。", ""),
        ("标签：说明", "一行里有多个单元格且只写了一个格式时，它们会合并成一个条目「第一格：其余各格」，例如「白灯（HL1）」+「市电电源指示灯」→「白灯（HL1）：市电电源指示灯」。", ""),
        ("A 列写正文", "如果某一行只有 A 列有文字、并且 J 列选了格式，那么 A 列的文字当作内容（不是标题）。", ""),
        ("表格", "把表格写在 B–I 列：表头行格式选「表头」（可以有多行），数据行选「表格」。合并单元格（包括跨行跨列）按原样合并。表格单元格里可以放图片。", ""),
        ("图片", "直接把图片插入到单元格。图片所在行如果是表格行，图片放进那个单元格；否则放在该行文字的下方。图片只支持 PNG / JPEG。", ""),
        ("列表编号", "连续的同类列表行是同一个列表；中间隔着「缩进段落」或图片时编号继续往下排；遇到标题、正文、表格后重新从 1 开始。", ""),
        ("", "", ""),
        ("上传后", "系统会立即解析，并在页面上预览章节结构；有问题会精确指出「工作表!单元格」。错误必须修正后才能生成，提醒可以确认后继续。", ""),
    ]
    for index, (a, b, c) in enumerate(rows, start=1):
        for col, value in enumerate((a, b, c), start=1):
            cell = sheet.cell(index, col, value)
            cell.alignment = wrap
        if a in {"整体规则", "格式一览", "进阶写法"}:
            for col in (1, 2, 3):
                sheet.cell(index, col).font = bold
                sheet.cell(index, col).fill = PatternFill("solid", fgColor="D9EAD3")
        elif a:
            sheet.cell(index, 1).font = bold
    sheet["A1"].font = title


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--case-dir", type=Path, default=ROOT / "SOP_case" / "升压一体机案例")
    parser.add_argument("--out", type=Path, default=ROOT / "templates" / "boost-integrated-unit" / "v1" / "input-template.xlsx")
    parser.add_argument("--report", action="store_true", help="print rows that could not be derived")
    args = parser.parse_args()

    source = args.case_dir / "升压一体机测试项分类表.xlsx"
    word = next(args.case_dir.glob("Glayx*.docx"))
    repaired = load_workbook(source)
    fix_cells(repaired)
    with tempfile.TemporaryDirectory() as scratch:
        staging = Path(scratch) / "repaired.xlsx"
        repaired.save(staging)  # openpyxl closes image streams on save, so analyse and write from separate loads
        grids, _ = read_workbook(staging)
        workbook = load_workbook(staging)
    chapter_grids = [grid for grid in grids if grid.name not in ("版本修改记录", "填写说明")]
    chapters = derive.read_word_chapters(word)
    unresolved = 0
    for grid, items in zip(chapter_grids, chapters, strict=True):
        annotations, report = derive.derive(grid, items, {})
        formats = {row: annotation.j for row, annotation in annotations.items() if annotation.j}
        formats.update(EXCEL_ONLY.get(grid.name, {}))
        sheet = workbook[grid.name]
        for row, value in formats.items():
            sheet.cell(row, FORMAT_COLUMN, value)
        style_and_validate(sheet)
        leftovers = [line for line in report if int(line.split("row ")[1].split(":")[0]) not in formats] if report else []
        unresolved += len(leftovers)
        if args.report:
            print(f"{grid.name}: {len(formats)} formatted rows" + (f", {len(leftovers)} unresolved" if leftovers else ""))
            for line in leftovers:
                print(line)
    add_instructions(workbook)
    workbook.save(args.out)
    print(f"saved {args.out} (unresolved rows: {unresolved})")


if __name__ == "__main__":
    main()
